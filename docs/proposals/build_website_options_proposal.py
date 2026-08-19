from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

# Writes the client-facing deliverable next to the other proposals, regardless
# of where the script is run from.
OUT = str(Path(__file__).resolve().parents[1] / "docs" / "proposals" / "Varniqa_Website_Expansion_Options.docx")
NAVY = "17324D"
TEAL = "2B7A78"
PALE = "EAF4F3"
LIGHT = "F4F6F8"
MID = "667580"
WHITE = "FFFFFF"
BLACK = "20262B"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc.get_or_add_tcPr()
    tcMar = tc.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tc.append(tcMar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")

def set_table_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    total = sum(widths)
    tblW.set(qn("w:w"), str(total)); tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.first_child_found_in("w:tblInd")
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd"); tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "140"); tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcW = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tcW is None:
                tcW = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(width)); tcW.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)

def font(run, size=10.5, bold=False, color=BLACK, italic=False):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

def para(doc, text="", size=10.5, bold=False, color=BLACK, after=7, before=0, align=None, italic=False):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.2
    if align is not None: p.alignment = align
    font(p.add_run(text), size, bold, color, italic)
    return p

def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.15
    font(p.add_run(text), 10.2)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.68); sec.left_margin = Inches(.85); sec.right_margin = Inches(.85)
sec.header_distance = Inches(.3); sec.footer_distance = Inches(.3)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(7); normal.paragraph_format.line_spacing = 1.2
for name, size, color, before, after in [("Heading 1", 17, NAVY, 15, 7), ("Heading 2", 12.5, TEAL, 10, 4)]:
    st = styles[name]; st.font.name = "Aptos Display"; st.font.size = Pt(size); st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color); st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
for sname in ("List Bullet", "List Bullet 2"):
    styles[sname].font.name = "Aptos"; styles[sname].font.size = Pt(10.2)

# Running header and footer
h = sec.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(h.add_run("VARNIQA SKIN SCIENCES  |  WEBSITE EXPANSION"), 8.5, True, MID)
f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(f.add_run("Website structure proposal  |  August 2026"), 8, False, MID)

# Cover/title block: proposal_centerpiece pattern
para(doc, "VARNIQA SKIN SCIENCES", 10, True, TEAL, 10, 12, WD_ALIGN_PARAGRAPH.CENTER)
para(doc, "Website Expansion Options", 27, True, NAVY, 5, 0, WD_ALIGN_PARAGRAPH.CENTER)
para(doc, "A clearer path from a single-page website to stronger organic search visibility", 12.5, False, MID, 20, 0, WD_ALIGN_PARAGRAPH.CENTER)

t = doc.add_table(rows=1, cols=2)
t.style = "Table Grid"
t.cell(0,0).text = "OPTION 1\nComplete multi-page website"
t.cell(0,1).text = "OPTION 2\nEssential SEO expansion"
set_table_widths(t, [4680, 4680])
for i, c in enumerate(t.rows[0].cells):
    shade(c, NAVY if i == 0 else TEAL)
    for j, p in enumerate(c.paragraphs):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
        for r in p.runs: font(r, 11 if j == 0 else 10, True, WHITE)

heading(doc, "Why the website structure matters", 1)
para(doc, "The current website presents the clinic, services, FAQs and contact details on one URL. This is useful for a quick overview, but it limits the number of focused pages that Google can understand, index and match to specific searches.")
para(doc, "A multi-page structure creates dedicated destinations for treatment, location and informational searches. It also improves internal linking, makes future content easier to publish, and gives prospective patients a clearer journey from discovery to appointment.")

call = doc.add_table(rows=1, cols=1); call.style = "Table Grid"; set_table_widths(call, [9360]); shade(call.cell(0,0), PALE)
p = call.cell(0,0).paragraphs[0]; p.paragraph_format.space_after = Pt(0)
font(p.add_run("Important: "), 10.2, True, TEAL); font(p.add_run("A better structure creates stronger SEO potential, but rankings also depend on content quality, competition, local authority, technical performance and ongoing optimisation."), 10.2)

doc.add_page_break()
heading(doc, "Option 1: Complete multi-page website", 1)
para(doc, "Best for long-term organic growth, a premium patient experience and a scalable digital presence.", 11, True, TEAL, 8)

heading(doc, "What is included", 2)
for x in [
    "Complete website relayout with a new navigation system and page hierarchy.",
    "Dedicated page for each priority service or treatment category, with focused copy, FAQs and clear calls to action.",
    "About the clinic and dermatologist pages to strengthen expertise, trust and credibility.",
    "Blog hub plus individual article pages for ongoing educational content.",
    "Dedicated contact page with location, hours, map and enquiry options.",
    "Dedicated appointment-booking page designed around conversion.",
    "SEO-supporting pages and elements, such as location-focused content, treatment FAQs, internal links, metadata, schema, breadcrumbs and sitemap updates where appropriate.",
    "Mobile-first layouts and clear pathways between services, educational content and booking."
]: bullet(doc, x)

heading(doc, "SEO and business value", 2)
for x in [
    "Each core service can target its own search intent instead of competing for space on the home page.",
    "The site can grow into a topic cluster: service pages answer commercial intent while blogs answer research questions and link patients to relevant care.",
    "More indexable, useful pages create more opportunities to appear for specific and local long-tail searches.",
    "Dedicated appointment and contact journeys reduce friction for patients who are ready to act.",
    "The structure can expand over time without repeatedly redesigning the main page."
]: bullet(doc, x)

heading(doc, "Considerations", 2)
para(doc, "This option requires a larger initial investment, more content planning and more client inputs or approvals. It delivers the strongest foundation, but results should be evaluated over time through indexing, rankings, qualified traffic and appointment enquiries.")

doc.add_page_break()
heading(doc, "Option 2: Essential SEO expansion", 1)
para(doc, "Best for a smaller immediate scope that improves structure while retaining most of the existing single-page experience.", 11, True, TEAL, 8)

heading(doc, "What is included", 2)
for x in [
    "Retain the existing home page and its overall layout.",
    "Add one separate Services page that organises all treatments or service categories.",
    "Add a Blog listing page with a basic structure for publishing articles.",
    "Update the primary navigation and add basic internal links between the home page, Services and Blog.",
    "Apply essential on-page and technical SEO foundations to the new pages, including page titles, descriptions, headings, clean URLs and sitemap updates."
]: bullet(doc, x)

heading(doc, "SEO and business value", 2)
for x in [
    "Creates additional indexable sections without a full redesign.",
    "Provides a starting point for publishing educational content.",
    "Improves content organisation and gives visitors a clearer overview of services.",
    "Requires less time, content and initial investment than Option 1."
]: bullet(doc, x)

heading(doc, "Limitations", 2)
for x in [
    "A combined Services page must cover many treatments, so it cannot target each service as precisely as individual pages.",
    "Most commercial keywords still depend on a small number of URLs.",
    "The appointment and contact experience remains less focused unless those pages are added later.",
    "Future growth may require another restructuring phase as the service and blog libraries expand."
]: bullet(doc, x)

doc.add_page_break()
heading(doc, "Side-by-side comparison", 1)
rows = [
    ("Website architecture", "Full multi-page structure", "Existing home + Services + Blog"),
    ("Service targeting", "Individual priority service pages", "All services grouped on one page"),
    ("Blog capability", "Blog hub and scalable article templates", "Basic blog listing and article structure"),
    ("Patient journey", "Dedicated contact and booking pages", "Primarily existing contact/booking flow"),
    ("SEO potential", "Highest: broader keyword and search-intent coverage", "Moderate: improved, but fewer focused URLs"),
    ("Scalability", "Designed for continued content and service growth", "Useful starting point; may need later expansion"),
    ("Initial effort", "Higher", "Lower"),
    ("Best fit", "Long-term growth and stronger market positioning", "Budget-conscious first phase")
]
tbl = doc.add_table(rows=1, cols=3); tbl.style = "Table Grid"
for i, txt in enumerate(["Decision area", "Option 1", "Option 2"]):
    c=tbl.rows[0].cells[i]; c.text=txt; shade(c, NAVY)
    for r in c.paragraphs[0].runs: font(r, 9.5, True, WHITE)
for label, a, b in rows:
    cells=tbl.add_row().cells
    for i, txt in enumerate((label,a,b)):
        cells[i].text=txt
        if i==0: shade(cells[i], LIGHT)
        for r in cells[i].paragraphs[0].runs: font(r, 9.2, i==0, NAVY if i==0 else BLACK)
set_table_widths(tbl, [2100, 3630, 3630])

heading(doc, "Recommendation", 1)
para(doc, "Option 1 is recommended if organic search is expected to become a meaningful patient-acquisition channel. It provides the page depth, topical structure and conversion paths needed for sustained SEO work.")
para(doc, "Option 2 is a practical phased alternative when budget or turnaround time is the immediate constraint. It should be presented as the first stage of a future multi-page build, not as an equivalent SEO outcome.")

rec = doc.add_table(rows=1, cols=1); rec.style="Table Grid"; set_table_widths(rec,[9360]); shade(rec.cell(0,0), PALE)
p=rec.cell(0,0).paragraphs[0]; font(p.add_run("Suggested decision rule\n"), 11, True, TEAL)
font(p.add_run("Choose Option 1 for the strongest long-term foundation. Choose Option 2 when the priority is a lower-cost first step, with a planned roadmap to add individual service, contact and appointment pages later."), 10.2)

heading(doc, "Proposed next step", 1)
para(doc, "Once an option is selected, the next phase is to confirm priority services, target locations, required clinic inputs, content responsibilities, timeline and commercial estimate. A keyword and competitor review should then guide the final page map and publishing priorities.")

doc.core_properties.title = "Varniqa Skin Sciences - Website Expansion Options"
doc.core_properties.subject = "Client proposal comparing two multi-page website approaches"
doc.core_properties.author = "Varniqa Website Team"
doc.save(OUT)
print(OUT)
